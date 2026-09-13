"""Parser tests — realistic Indian bank / card / Tally exports built in memory."""
import io
from datetime import date
import pytest
from openpyxl import Workbook

from app.parsers.tabular import parse_rows, _date, _money
from app.parsers.tally_xml import parse_tally_xml


def xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def by_desc(rows):
    return {r["raw_description"]: r for r in rows}


# --- cell parsing ------------------------------------------------------------------------
def test_date_shapes():
    assert _date("2026-04-01 00:00:00") == date(2026, 4, 1)      # ISO is never day-first
    assert _date("20260401") == date(2026, 4, 1)                 # Tally yyyymmdd
    assert _date("01/04/26") == date(2026, 4, 1)
    assert _date("1 Apr 2026") == date(2026, 4, 1)               # SBI
    assert _date("1-Apr-26") == date(2026, 4, 1)                 # Tally
    assert _date("01/04/2026 10:35 PM") == date(2026, 4, 1)
    assert _date("45000") == date(2023, 3, 15)                   # Excel serial
    for junk in ("1", "12", "Page 1 of 3", "MARCH 2026", "TOTAL", "Opening Balance", "31/02/2026"):
        assert _date(junk) is None, junk


def test_money_shapes():
    assert _money("1,45,000.00") == 145000.0
    assert _money("Rs. 5,310.00") == 5310.0
    assert _money("1,769.00 Dr") == 1769.0
    assert _money("(500.00)") == -500.0
    assert _money("45,000.00-") == -45000.0
    assert _money("-") is None and _money(" ") is None and _money("nan") is None


# --- bank formats ------------------------------------------------------------------------
def test_icici_csv_preamble_and_legends():
    csv = (
        "Account Statement,,,,,,,\nCustomer Name,ACME,,,,,,\n,,,,,,,\n,,,,,,,\n"
        "S No.,Value Date,Transaction Date,Cheque Number,Transaction Remarks,Withdrawal Amount (INR ),Deposit Amount (INR ),Balance (INR )\n"
        '1,03/04/2026,03/04/2026,,VIN/AMAZON INTERNET SERVICES PVT LTD/20260403123456/ECOM,"9,440.00",0.00,"5,40,560.00"\n'
        '2,15/04/2026,15/04/2026,,NEFT-CUSTOMER PAYMENT-JINDAL STEEL,0.00,"1,20,000.00","6,33,740.00"\n'
        ",,,,,,,\nLegends,,,,,,,\nVIN,Visa International,,,,,,\n"
    ).encode()
    rows, fmt, skipped = parse_rows(csv, "icici.csv", "bank")
    assert fmt == "bank_debit_credit" and len(rows) == 2
    assert rows[0]["amount"] == 9440 and rows[0]["direction"] == "debit"
    assert rows[1]["amount"] == 120000 and rows[1]["direction"] == "credit"


def test_kotak_csv_narrow_preamble_wider_table():
    """The first line has fewer commas than the table: every table row must survive."""
    csv = (
        "Kotak Mahindra Bank,,,,,\nAccount Statement,,,,,\n"
        "Sl. No.,Date,Narration,Chq/Ref No,Withdrawal (Dr),Deposit(Cr),Balance\n"
        '1,01-04-2026,UPI/tallysolutions.rzp@hdfcbank/609112345678/TSS,UPI-609112345678,"5,310.00",,"1,00,000.00"\n'
        '2,05-04-2026,IMPS 610512345 CUSTOMER LTD,IMPS,,"10,000.00","87,688.00"\n'
    ).encode()
    rows, _, _ = parse_rows(csv, "kotak.csv", "bank")
    assert len(rows) == 2
    assert rows[0]["ref"] == "UPI-609112345678" and rows[0]["direction"] == "debit"
    assert rows[1]["direction"] == "credit"


def test_sbi_xlsx_text_dates_and_blank_ref():
    data = [
        ["Account Name", ":", "ACME"], [],
        ["Txn Date", "Value Date", "Description", "Ref No./Cheque No.", "        Debit", "Credit", "Balance"],
        ["1 Apr 2026", "1 Apr 2026", "TO TRANSFER-UPI/DR/109876543210/GODADDY/ICIC/godaddy.rzp@icici/Pay--", "TRANSFER TO 4897", "1,299.00", " ", "3,45,000.00"],
        ["8 Apr 2026", "8 Apr 2026", "POS 4XXXXXXXXXXX1234 ZOOM.US", "", "1,769.00", "", "3,84,381.00"],
        ["", "", "**This is a computer generated statement", "", "", "", ""],
    ]
    rows, _, _ = parse_rows(xlsx(data), "sbi.xlsx", "bank")
    assert [r["date"] for r in rows] == [date(2026, 4, 1), date(2026, 4, 8)]
    assert rows[1]["ref"] is None          # never the string "nan"


def test_axis_xlsx_totals_row_skipped():
    data = [
        ["Name :- ACME"], [],
        ["Tran Date", "CHQNO", "PARTICULARS", "DR", "CR", "BAL", "SOL"],
        ["01-04-2026", "", "ECOM PUR/GOOGLE *GSUITE_acme.in/010426/1234", 2360, None, 100000, "1234"],
        ["06-04-2026", "", "IMPS/P2A/610612345678/CUSTOMER", None, 25000, 77800, "1234"],
        ["", "", "TOTAL", 2360, 25000, "", ""],
    ]
    rows, _, skipped = parse_rows(xlsx(data), "axis.xlsx", "bank")
    assert len(rows) == 2 and {r["direction"] for r in rows} == {"debit", "credit"}


def test_hdfc_xlsx_merged_header_blank_rows_summary_and_dotted_ref_column():
    wb = Workbook()
    ws = wb.active
    ws.append(["HDFC BANK LTD"]); ws.merge_cells("A1:G1")
    ws.append(["Statement of account"]); ws.merge_cells("A2:G2")
    ws.append([])
    ws.append(["Date", "Narration", "Chq./Ref.No.", "Value Dt", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"])
    ws.append(["01/04/26", "POS 412345XXXXXX1234 MSFT * E0100ABCD REDMOND WA", "0000000000000123", "01/04/26", "1,450.00", "", "1,45,000.00"])
    ws.append([])
    ws.append(["06/04/26", "NEFT CR-HDFC0000001-CUSTOMER-ACME-N096261234", "N096261234", "06/04/26", "", "1,00,000.00", "2,30,575.00"])
    ws.append([])
    ws.append(["", "STATEMENT SUMMARY :-", "", "", "", "", ""])
    ws.append(["", "", "", "", "15,875.00", "1,00,000.00", ""])
    b = io.BytesIO(); wb.save(b)
    rows, _, _ = parse_rows(b.getvalue(), "hdfc.xlsx", "bank")
    assert len(rows) == 2
    assert rows[0]["ref"] == "0000000000000123" and rows[0]["amount"] == 1450 and rows[0]["date"] == date(2026, 4, 1)


def test_credit_card_single_amount_with_dr_cr_suffix():
    csv = (
        "HDFC Bank Credit Card Statement,,,\nCard No: 4123 XXXX XXXX 5678,,,\n,,,\n"
        "Date,Transaction Description,Amount,Reward Points\n"
        '03/04/2026,ZOOM.US 888-799-9666 SAN JOSE,"1,769.00 Dr",17\n'
        '10/04/2026,PAYMENT RECEIVED - THANK YOU,"12,000.00 Cr",0\n'
        '12/04/2026,ZOOM.US REFUND,"1,769.00 Cr",0\n'
    ).encode()
    rows, fmt, _ = parse_rows(csv, "card.csv", "card")
    assert fmt == "bank_single_amount"
    d = by_desc(rows)
    assert d["ZOOM.US 888-799-9666 SAN JOSE"]["direction"] == "debit"
    assert d["PAYMENT RECEIVED - THANK YOU"]["direction"] == "credit"
    assert d["ZOOM.US REFUND"]["direction"] == "credit" and d["ZOOM.US REFUND"]["amount"] == 1769


def test_empty_and_header_only_files():
    with pytest.raises(ValueError):
        parse_rows(b"", "empty.csv", "bank")
    rows, _, _ = parse_rows(b"Date,Narration,Withdrawal Amt.,Deposit Amt.\n", "hdr.csv", "bank")
    assert rows == []


def test_legacy_xls():
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")
    for i, r in enumerate([["Bank of Baroda"], [], ["Txn Date", "Description", "Debit", "Credit", "Balance"],
                           ["02/04/2026", "NEFT-N1-TALLY SOLUTIONS PVT LTD-TSS", "5310.00", "", "50000"]]):
        for j, v in enumerate(r):
            ws.write(i, j, v)
    b = io.BytesIO(); wb.save(b)
    rows, _, _ = parse_rows(b.getvalue(), "bob.xls", "bank")
    assert len(rows) == 1 and rows[0]["amount"] == 5310


def test_pdf_table():
    pytest.importorskip("pdfplumber")
    rl = pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    from reportlab.lib import colors
    b = io.BytesIO()
    t = Table([["Date", "Narration", "Withdrawal", "Deposit", "Balance"],
               ["01/04/2026", "POS MSFT * E0100ABCD", "1,450.00", "", "1,00,000.00"],
               ["05/04/2026", "IMPS CR CUSTOMER", "", "20,000.00", "1,14,690.00"]])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(b, pagesize=A4).build([t])
    rows, _, _ = parse_rows(b.getvalue(), "stmt.pdf", "bank")
    assert len(rows) == 2 and rows[0]["amount"] == 1450 and rows[1]["direction"] == "credit"


def test_5000_rows_under_5s():
    import time
    lines = ["HDFC BANK,,,,", "Date,Narration,Chq./Ref.No.,Withdrawal Amt.,Deposit Amt.,Closing Balance"]
    for i in range(5000):
        lines.append(f"{1 + i % 28:02d}/{1 + i % 12:02d}/25,POS VENDOR {i % 40},{i},\"{(i % 50) * 137 + 1000:,}.00\",,\"1,00,000.00\"")
    t = time.time()
    rows, _, _ = parse_rows("\n".join(lines).encode(), "big.csv", "bank")
    assert len(rows) == 5000 and time.time() - t < 5


# --- Tally --------------------------------------------------------------------------------
def test_tally_day_book_ledger_lines_direction_and_payment_dedupe():
    data = [
        ["Day Book"], ["ACME"], [],
        ["Date", "Particulars", "Vch Type", "Vch No.", "Debit", "Credit"],
        ["1-Apr-26", "Cadspro Technologies Pvt Ltd", "Purchase", "12", None, "47,200.00"],
        [None, "SolidWorks Subscription Service", None, None, "40,000.00", None],
        [None, "Input CGST", None, None, "3,600.00", None],
        [None, "Input SGST", None, None, "3,600.00", None],
        ["3-Apr-26", "HDFC Bank", "Payment", "45", None, "46,400.00"],       # settles the purchase above -> dropped
        [None, "Cadspro Technologies Pvt Ltd", None, None, "46,400.00", None],
        ["5-Apr-26", "Cash", "Payment", "46", None, "8,850.00"],            # cash payment, no purchase voucher -> kept
        [None, "ABC Computers", None, None, "8,850.00", None],
        ["7-Apr-26", "Jindal Steel", "Receipt", "3", "1,20,000.00", None],
        ["9-Apr-26", "Customer", "Sales", "7", "50,000.00", None],
    ]
    rows, fmt, _ = parse_rows(xlsx(data), "daybook.xlsx", "tally")
    assert fmt == "tally_register"
    d = by_desc(rows)
    sw = d["Cadspro Technologies Pvt Ltd | SolidWorks Subscription Service"]
    assert sw["direction"] == "debit" and sw["amount"] == 47200 and sw["gst"] == 7200 and sw["taxable"] == 40000
    assert sw["ledger"] == "SolidWorks Subscription Service"
    assert d["ABC Computers"]["amount"] == 8850 and d["ABC Computers"]["direction"] == "debit"
    assert d["Jindal Steel"]["direction"] == "credit"
    assert not any("HDFC Bank" in k or "Customer" in k for k in d)


def test_tally_purchase_register_columnar_still_works():
    data = [["Date", "Particulars", "Vch Type", "Vch No.", "GSTIN/UIN", "Taxable Value", "GST", "Gross Total", "Narration"],
            ["01-Mar-26", "Cadspro Technologies Pvt Ltd", "Purchase", "SW2026", "36AAACC1234A1Z5", 40000, 7200, 47200, "SolidWorks Subscription Service 1 seat"]]
    rows, fmt, _ = parse_rows(xlsx(data), "register.xlsx", "tally")
    assert fmt == "tally_register" and len(rows) == 1
    r = rows[0]
    assert r["amount"] == 47200 and r["taxable"] == 40000 and r["gst"] == 7200 and r["direction"] == "debit"
    assert r["raw_description"].startswith("Cadspro Technologies Pvt Ltd | SolidWorks")


def test_tally_xml_payment_voucher_party_and_dedupe():
    xml = b"""<ENVELOPE><BODY><DATA><TALLYMESSAGE>
<VOUCHER VCHTYPE="Purchase"><DATE>20260401</DATE><PARTYLEDGERNAME>Cadspro Technologies Pvt Ltd</PARTYLEDGERNAME><VOUCHERNUMBER>12</VOUCHERNUMBER><NARRATION>SolidWorks &amp; support</NARRATION>
<ALLLEDGERENTRIES.LIST><LEDGERNAME>Cadspro Technologies Pvt Ltd</LEDGERNAME><AMOUNT>47200.00</AMOUNT></ALLLEDGERENTRIES.LIST>
<ALLLEDGERENTRIES.LIST><LEDGERNAME>Software Subscription</LEDGERNAME><AMOUNT>-40000.00</AMOUNT></ALLLEDGERENTRIES.LIST>
<ALLLEDGERENTRIES.LIST><LEDGERNAME>Input CGST</LEDGERNAME><AMOUNT>-3600.00</AMOUNT></ALLLEDGERENTRIES.LIST>
<ALLLEDGERENTRIES.LIST><LEDGERNAME>Input SGST</LEDGERNAME><AMOUNT>-3600.00</AMOUNT></ALLLEDGERENTRIES.LIST>
</VOUCHER>
<VOUCHER VCHTYPE="Payment"><DATE>20260403</DATE><PARTYLEDGERNAME>HDFC Bank</PARTYLEDGERNAME><VOUCHERNUMBER>45</VOUCHERNUMBER>
<ALLLEDGERENTRIES.LIST><LEDGERNAME>Cadspro Technologies Pvt Ltd</LEDGERNAME><AMOUNT>-46400.00</AMOUNT></ALLLEDGERENTRIES.LIST>
<ALLLEDGERENTRIES.LIST><LEDGERNAME>HDFC Bank</LEDGERNAME><AMOUNT>46400.00</AMOUNT></ALLLEDGERENTRIES.LIST>
</VOUCHER>
<VOUCHER VCHTYPE="Payment"><DATE>20260405</DATE><PARTYLEDGERNAME>Cash</PARTYLEDGERNAME><VOUCHERNUMBER>46</VOUCHERNUMBER><NARRATION>AMC Q1</NARRATION>
<ALLLEDGERENTRIES.LIST><LEDGERNAME>ABC Computers</LEDGERNAME><AMOUNT>-8850.00</AMOUNT></ALLLEDGERENTRIES.LIST>
<ALLLEDGERENTRIES.LIST><LEDGERNAME>Cash</LEDGERNAME><AMOUNT>8850.00</AMOUNT></ALLLEDGERENTRIES.LIST>
</VOUCHER>
</TALLYMESSAGE></DATA></BODY></ENVELOPE>"""
    rows, fmt, _ = parse_tally_xml(xml)
    assert fmt == "tally_xml" and len(rows) == 2
    assert rows[0]["amount"] == 47200 and rows[0]["gst"] == 7200
    assert rows[1]["raw_description"] == "ABC Computers | AMC Q1" and rows[1]["amount"] == 8850
