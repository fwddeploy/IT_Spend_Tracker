"""Excel export: Lines, Upcoming (12 months), Payments, Questions — openpyxl, one workbook in memory."""
from __future__ import annotations
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models as m


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([v.isoformat() if isinstance(v, date) else v for v in r])
    for i, h in enumerate(header, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None]) if rows else len(h)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    return ws


def build_workbook(db: Session, company: m.Company, streams_out: list[dict], upcoming: list[dict]) -> bytes:
    wb = Workbook()
    _sheet(wb, "Lines",
           ["Vendor", "Product", "Category", "Type", "Cycle", "Expected amount", "Monthly equivalent", "Last paid", "Next due", "Status",
            "Auto-renew", "Paid from", "Owner", "Pay URL", "Reminder", "Confidence", "Payments seen", "Sources", "Flags", "Notes"],
           [[s["vendor_name"], s["product"], s["category"], s["stream_type"], s["cycle"], s["expected_amount"], s["monthly_equivalent"],
             s["last_paid_date"], s["next_due"], s["status"], "yes" if s["auto_renew"] else "no", s["paid_from"], s["owner_name"],
             s["pay_url"], "on" if s["reminder_on"] else "off", s["confidence"], s["occurrences_count"], ", ".join(s["sources"]),
             "; ".join(s.get("flags_human") or []), s["notes"]] for s in streams_out], first=True)
    up_rows = []
    for mth in upcoming:
        for it in mth["items"]:
            up_rows.append([mth["month"], it["due"], it["vendor_name"], it["product"], it["expected_amount"], it["status"], it["paid_from"]])
    _sheet(wb, "Upcoming", ["Month", "Due", "Vendor", "Product", "Amount", "Status", "Paid from"], up_rows)
    sid2name = {s["id"]: (s["vendor_name"], s["product"]) for s in streams_out}
    occs = db.execute(select(m.Occurrence).where(m.Occurrence.company_id == company.id).order_by(m.Occurrence.date.desc())).scalars().all()
    _sheet(wb, "Payments", ["Date", "Vendor", "Product", "Payee", "Paid", "Gross", "Taxable", "GST", "TDS", "Fees", "Currency", "Foreign amount",
                            "Mode", "Invoice", "Period from", "Period to", "Sources", "Unpaid", "Narration"],
           [[o.date, sid2name.get(o.stream_id, ("", ""))[0], o.product or sid2name.get(o.stream_id, ("", ""))[1], o.payee_clean, o.amount_paid,
             o.amount_gross, o.taxable, o.gst, o.tds, o.fees, o.currency, o.fx_amount, o.payment_mode, o.invoice_no, o.period_from, o.period_to,
             ", ".join(o.sources or []), "yes" if o.unpaid else "no", o.raw_description] for o in occs])
    qs = db.execute(select(m.Question).where(m.Question.company_id == company.id).order_by(m.Question.id)).scalars().all()
    _sheet(wb, "Questions", ["Asked", "Kind", "Question", "Answered", "Answer"],
           [[q.created_at.date(), q.kind, q.prompt, "yes" if q.answered else "no", q.answer] for q in qs])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
